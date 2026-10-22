"""Build data/eagle_vocabulary.json from the official EAGLE matrix workbooks.

The EAGLE Barcoding Template (v3.2) provides URI, code, label and level of every
matrix element; the 3.2 change-log workbook adds the element type (Heading /
Parameter / Character) and LC/LU relevance for each element.

Usage:
    uv run --group build python scripts/build_vocabulary.py [BARCODING_DIR]
"""

import json
import sys
from pathlib import Path

import openpyxl

DEFAULT_DIR = Path(
    r"C:\Users\rosenauer\OneDrive\OneDrive - Umweltbundesamt GmbH\Projekte\EAGLE"
    r"\Task 8_ Use of AI\Barcoding"
)
TEMPLATE = "EAGLE Barcoding Template_3.2_20230829.xlsx"
CHANGELOG = "EAGLE Matrix ChangeLogFile Old3.1.2-New3.2_20240606_corr.xlsx"
OUT = Path(__file__).resolve().parent.parent / "data" / "eagle_vocabulary.json"

BLOCKS = {
    "LCC": "Land Cover Components",
    "LUA": "Land Use Attributes",
    "LCH": "Land Characteristics",
}


def parent_code(code: str) -> str | None:
    if "_" in code:
        return code.rsplit("_", 1)[0]
    if "-" in code:
        return code.split("-", 1)[0]
    return None


def read_metadata(path: Path) -> dict[str, tuple[str | None, str | None]]:
    wb = openpyxl.load_workbook(path, read_only=True)
    meta = {}
    for row in wb["Matrix_Vertical_CLEAN_NEW"].iter_rows(min_row=2, values_only=True):
        if row[1]:
            meta[row[1]] = (row[4], row[5])
    return meta


def main(src: Path) -> None:
    meta = read_metadata(src / CHANGELOG)
    wb = openpyxl.load_workbook(src / TEMPLATE, read_only=True)
    concepts = []
    for row in wb["Matrix_Vertical_BarCoding"].iter_rows(min_row=6, values_only=True):
        uri, code, label, level = row[:4]
        if not code:
            continue
        relevance, element_type = meta.get(code, (None, None))
        concepts.append(
            {
                "code": code,
                "uri": uri,
                "label": label.strip(),
                "level": int(level),
                "block": code.split("-")[0],
                "parent": parent_code(code),
                "element_type": element_type,
                "relevance": relevance,
            }
        )
    payload = {
        "source": TEMPLATE,
        "version": "3.2",
        "blocks": BLOCKS,
        "concepts": concepts,
    }
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps(payload, indent=1, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {len(concepts)} concepts to {OUT}")


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_DIR)
