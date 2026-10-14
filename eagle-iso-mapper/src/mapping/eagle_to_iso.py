import io
import re
import xml.etree.ElementTree as ET
from xml.dom import minidom
from typing import List, Optional, Dict, Any
import openpyxl

from src.python.generated.eagle_iso_models import (
    LandUnit,
    LandCoverComponent,
    AbioticComponent,
    BioticComponent,
    WaterComponent,
    LandUseAttribute,
    LandCharacteristic,
    BarCodeValueEnum,
    ISOPresenceEnum,
    FunctionalRelationshipEnum,
    TemporalRelationshipEnum
)
from src.mapping.sssom_engine import sssom_resolver


class EagleXslxParser:
    """
    Parses vertical EAGLE Matrix Barcoding sheets into validated LandUnit Pydantic objects.
    Uses SSSOM crosswalk for semantic bridging between EAGLE codes/labels and ISO concepts.
    """

    def __init__(self, workbook: openpyxl.Workbook, sheet_name: str, log_callback=print):
        self.workbook = workbook
        self.sheet = self.workbook[sheet_name]
        self.log = log_callback

    def extract_land_units(self) -> List[LandUnit]:
        self.log("Identifying Class Codes and Names from sheet...")
        classes_info: Dict[int, Dict[str, Any]] = {}

        # 1. Locate header rows dynamically
        code_row = 1
        for r in range(1, 10):
            val = self.sheet.cell(row=r, column=1).value
            if val and "Class Code" in str(val):
                code_row = r
                break

        name_row = code_row + 1

        # 2. Identify Class columns (starting at col 5, step 2)
        for col_idx in range(5, self.sheet.max_column + 1, 2):
            class_code = self.sheet.cell(row=code_row, column=col_idx).value
            class_name = self.sheet.cell(row=name_row, column=col_idx).value

            if class_code and str(class_code).strip() != "":
                clean_code = str(class_code).strip()
                clean_name = str(class_name).strip() if class_name else clean_code
                classes_info[col_idx] = {
                    "code": clean_code,
                    "name": clean_name,
                    "land_cover_components": [],
                    "land_use_attributes": [],
                    "land_characteristics": []
                }

        if not classes_info:
            err_msg = f"No Land Unit classes found in sheet '{self.sheet.title}'. Please ensure the sheet has 'Class Code' headers (e.g. 'CLC-all 3 Classes' or 'CLC112')."
            self.log(f"⚠️ {err_msg}")
            raise ValueError(err_msg)

        self.log(f"Found {len(classes_info)} Land Unit classes in sheet. Parsing taxonomy rows...")

        # 3. Iterate over concept rows
        start_row = name_row + 4  # usually row 6 or 7
        for row_idx in range(start_row, self.sheet.max_row + 1):
            uri_val = self.sheet.cell(row=row_idx, column=1).value
            code_val = self.sheet.cell(row=row_idx, column=2).value
            label_val = self.sheet.cell(row=row_idx, column=3).value

            if not code_val and not label_val:
                continue

            code_str = str(code_val).strip() if code_val else ""
            label_str = str(label_val).strip() if label_val else ""
            uri_str = str(uri_val).strip() if uri_val else ""

            # Skip root segment headings
            if code_str.upper() in ["LCC", "LUA", "LCH"] or label_str.lower() in [
                "land cover components", "land use attributes", "land characteristics"
            ]:
                continue

            # Check each class column for barcode entries
            for col_idx, class_data in classes_info.items():
                barcode_cell = self.sheet.cell(row=row_idx, column=col_idx).value
                factor_cell = self.sheet.cell(row=row_idx, column=col_idx + 1).value

                if barcode_cell is None:
                    continue

                barcode_str = str(barcode_cell).strip()
                if not barcode_str or barcode_str.upper() in ["SEGMENT HEADING", "NONE", "LCC", "LUA", "LCH"]:
                    continue

                factor_str = str(factor_cell).strip() if factor_cell is not None else None

                # Resolve SSSOM crosswalk mapping
                sssom_match = sssom_resolver.resolve_eagle_to_iso(code_str, label_str)
                iso_presence = sssom_resolver.map_bcv_to_presence(barcode_str)
                cov_min, cov_max = sssom_resolver.factor_to_coverage(factor_str)

                bcv_enum = None
                bcv_key = f"BCV_{barcode_str.upper()}"
                if bcv_key in BarCodeValueEnum.__members__:
                    bcv_enum = BarCodeValueEnum[bcv_key]
                elif barcode_str in BarCodeValueEnum.__members__:
                    bcv_enum = BarCodeValueEnum[barcode_str]

                # Component categorization
                if code_str.startswith("LCC") or "land cover" in label_str.lower():
                    # Check for duplicates
                    already_exists = any(
                        c.uri == uri_str or (c.label and c.label == label_str)
                        for c in class_data["land_cover_components"]
                    )
                    if not already_exists:
                        # Instantiate specific subtype according to nature of component
                        is_abiotic = any(term in label_str.lower() for term in ["sealed", "surface", "soil", "rock", "construction", "building"])
                        is_water = any(term in label_str.lower() for term in ["water", "snow", "ice"])

                        comp_kwargs = dict(
                            uri=uri_str or None,
                            label=label_str,
                            bcv_operator=bcv_enum,
                            iso_presence_type=iso_presence,
                            min_cover=cov_min,
                            max_cover=cov_max,
                            stratum_id=1
                        )

                        if is_abiotic:
                            comp = AbioticComponent(**comp_kwargs)
                        elif is_water:
                            comp = WaterComponent(**comp_kwargs)
                        else:
                            comp = BioticComponent(**comp_kwargs)

                        class_data["land_cover_components"].append(comp)

                elif code_str.startswith("LUA") or "land use" in label_str.lower():
                    # Extract HILUCS code from code_str (e.g. LUA-5_1 -> 5.1)
                    hilucs_code = None
                    if code_str.startswith("LUA-"):
                        hilucs_code = code_str.replace("LUA-", "").replace("_", ".")

                    lua = LandUseAttribute(
                        uri=uri_str or None,
                        label=label_str,
                        hilucs_code=hilucs_code,
                        iso_function_type=sssom_match.get("object_label") if sssom_match else None,
                        functional_relationship=FunctionalRelationshipEnum.Primary if barcode_str in ["3", "5"] else FunctionalRelationshipEnum.Secondary,
                        temporal_relationship=TemporalRelationshipEnum.Coexistant
                    )
                    class_data["land_use_attributes"].append(lua)

                elif code_str.startswith("LCH") or "land characteristic" in label_str.lower():
                    char_type = "Context"
                    if "irrigation" in label_str.lower():
                        char_type = "Irrigation"
                    elif any(s in label_str.lower() for s in ["rice", "hops", "species"]):
                        char_type = "FloristicAspect"

                    lch = LandCharacteristic(
                        uri=uri_str or None,
                        label=label_str,
                        characteristic_type=char_type
                    )
                    class_data["land_characteristics"].append(lch)

        # 4. Construct validated LandUnit Pydantic objects
        validated_units: List[LandUnit] = []
        for col_idx, data in classes_info.items():
            unit = LandUnit(
                id=data["code"],
                name=data["name"],
                geometry_type="Polygon",
                land_cover_components=data["land_cover_components"],
                land_use_attributes=data["land_use_attributes"],
                land_characteristics=data["land_characteristics"]
            )
            validated_units.append(unit)

        self.log(f"Successfully constructed {len(validated_units)} validated LandUnits.")
        return validated_units


class LchsXmlGenerator:
    """
    Generates standard ISO 19144 (LChS LCML/LUML) XML from validated LandUnit objects.
    Produces the standard relational structure supported by ISO 19144 tools.
    """

    def __init__(self, log_callback=print):
        self.log = log_callback

    def _prettify(self, elem: ET.Element) -> bytes:
        rough_string = ET.tostring(elem, encoding="utf-8")
        reparsed = minidom.parseString(rough_string)
        return reparsed.toprettyxml(indent="  ", encoding="utf-8")

    def generate_xml(self, land_units: List[LandUnit]) -> bytes:
        root = ET.Element(
            "LC_Legend",
            attrib={
                "id": "1",
                "legend_name": "EAGLE to ISO 19144 Crosswalk Legend",
                "legend_description": "Exported from EAGLE Matrix Barcoding via SSSOM",
                "legend_author": "EAGLE ISO Mapper",
                "elementID": "1",
                "formElementID": "1",
                "objectID": "1",
                "objectReference": "LC_Legend"
            }
        )
        objects = ET.SubElement(root, "objects")

        self.log(f"Generating ISO 19144 XML for {len(land_units)} Land Units...")

        # 1. Classes
        for idx, unit in enumerate(land_units, 1):
            c_elem = ET.SubElement(objects, "LC_Class")
            ET.SubElement(c_elem, "legend_id").text = "1"
            ET.SubElement(c_elem, "class_id").text = str(idx)
            ET.SubElement(c_elem, "class_name").text = unit.name or unit.id
            ET.SubElement(c_elem, "class_description").text = f"EAGLE Class {unit.id}"
            ET.SubElement(c_elem, "class_map_code").text = unit.id

        # 2. Horizontal Patterns
        for idx in range(1, len(land_units) + 1):
            hp_elem = ET.SubElement(objects, "LC_HorizontalPatterns")
            ET.SubElement(hp_elem, "class_id").text = str(idx)
            ET.SubElement(hp_elem, "horizontal_pattern_id").text = str(idx)
            ET.SubElement(hp_elem, "name").text = f"Horizontal Pattern {idx}"
            ET.SubElement(hp_elem, "cover").text = "0"
            ET.SubElement(hp_elem, "cover").text = "100"
            ET.SubElement(hp_elem, "occurrence").text = "0"
            ET.SubElement(hp_elem, "occurrence").text = "100"

        # 3. Strata
        for idx in range(1, len(land_units) + 1):
            strata_elem = ET.SubElement(objects, "LC_Strata")
            ET.SubElement(strata_elem, "HPID").text = str(idx)
            ET.SubElement(strata_elem, "stratumID").text = str(idx)
            ET.SubElement(strata_elem, "name").text = f"Stratum {idx}"
            ET.SubElement(strata_elem, "presenceType").text = "Fixed"
            ET.SubElement(strata_elem, "portioning").text = "100"

        # 4. Properties (Land Cover Components)
        for idx, unit in enumerate(land_units, 1):
            if not unit.land_cover_components:
                continue

            for comp in unit.land_cover_components:
                sssom_match = sssom_resolver.resolve_eagle_to_iso(comp.uri, comp.label)
                if sssom_match:
                    obj_id = sssom_match.get("object_id", "")
                    block_ref = sssom_resolver._normalize_curie(obj_id)
                else:
                    # Fallback clean label
                    clean_name = re.sub(r'[^a-zA-Z0-9]', '', comp.label or "Vegetation")
                    block_ref = f"LC_{clean_name}"

                presence_str = getattr(comp.iso_presence_type, "value", str(comp.iso_presence_type)) if comp.iso_presence_type else "Fixed"

                prop_elem = ET.SubElement(objects, "LC_Properties")
                ET.SubElement(prop_elem, "StratumID").text = str(idx)
                ET.SubElement(prop_elem, "BlockReference").text = block_ref
                ET.SubElement(prop_elem, "elementPresenceType").text = presence_str

                c_min = comp.min_cover if comp.min_cover is not None else 0
                c_max = comp.max_cover if comp.max_cover is not None else 100
                ET.SubElement(prop_elem, "cover").text = str(c_min)
                ET.SubElement(prop_elem, "cover").text = str(c_max)

        # 5. Characteristics (Land Characteristics)
        for idx, unit in enumerate(land_units, 1):
            if not unit.land_characteristics:
                continue

            for char in unit.land_characteristics:
                lbl = (char.label or "").lower()
                if "irrigation" in lbl:
                    irr_elem = ET.SubElement(objects, "LC_Characteristics")
                    ET.SubElement(irr_elem, "StratumID").text = str(idx)
                    ET.SubElement(irr_elem, "BlockReference").text = "LC_Graminoid"
                    ET.SubElement(irr_elem, "CharacteristicReference").text = "LC_CultivatedAndManagedVegetationCharacteristics"
                    ET.SubElement(irr_elem, "CharacteristicLabel").text = "Cultivated & Managed Vegetation"
                    irr_type = "Sprinkler" if "sprinkler" in lbl else "Unspecified"
                    ET.SubElement(irr_elem, "irrigationType").text = irr_type

                elif any(sp in lbl for sp in ["rice", "hops"]):
                    sp_elem = ET.SubElement(objects, "LC_Characteristics")
                    ET.SubElement(sp_elem, "StratumID").text = str(idx)
                    ET.SubElement(sp_elem, "BlockReference").text = "LC_Graminoid"
                    ET.SubElement(sp_elem, "CharacteristicReference").text = "LC_FloristicAspectsCharacteristic"
                    ET.SubElement(sp_elem, "CharacteristicLabel").text = "Floristic Aspects"
                    species = "Oryza sativa" if "rice" in lbl else "Humulus lupulus"
                    ET.SubElement(sp_elem, "speciesName").text = species

        return self._prettify(root)


def run_eagle_to_iso_mapping(excel_workbook: openpyxl.Workbook, sheet_name: str, log_callback=print) -> bytes:
    log_callback("Initializing EagleXslxParser with SSSOM Resolver...")
    parser = EagleXslxParser(excel_workbook, sheet_name, log_callback)

    log_callback("Extracting and semantically validating data into LandUnits...")
    land_units = parser.extract_land_units()

    log_callback(f"Successfully processed {len(land_units)} Land Units. Generating ISO 19144 XML...")
    generator = LchsXmlGenerator(log_callback)

    return generator.generate_xml(land_units)