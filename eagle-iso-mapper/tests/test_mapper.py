import io
import xml.etree.ElementTree as ET
import openpyxl
import pytest
from pathlib import Path

from src.mapping.sssom_engine import SSSOMResolver
from src.mapping.eagle_to_iso import run_eagle_to_iso_mapping
from src.mapping.iso_to_eagle import run_iso_to_eagle_mapping
from src.python.generated.eagle_iso_models import (
    LandUnit,
    LandCoverComponent,
    BarCodeValueEnum,
    ISOPresenceEnum
)

TEST_EXCEL_PATH = Path("testdata/eagle/EAGLE Barcoding Examples CLC - v4.xlsx")
TEST_XML_PATH = Path("testdata/iso/LChS_Legend.lchs.xml")


def test_sssom_resolver_load():
    resolver = SSSOMResolver()
    assert resolver.msdf is not None
    assert not resolver.msdf.df.empty
    assert len(resolver.msdf.df) >= 20

    # Test forward lookup
    sealed = resolver.resolve_eagle_to_iso("LCC-1_1_1")
    assert sealed is not None
    assert "LC_BuildUpSurface" in sealed["object_id"] or "LC_BuiltUpSurface" in sealed["object_id"]

    # Test reverse lookup
    rev_soil = resolver.resolve_iso_to_eagle("LC_BareSoil")
    assert rev_soil is not None
    assert "LCC-1_2_2_2" in rev_soil["subject_id"]


def test_sssom_bcv_and_coverage_mapping():
    resolver = SSSOMResolver()

    # Presence mappings
    assert resolver.map_bcv_to_presence("3") == ISOPresenceEnum.Fixed
    assert resolver.map_bcv_to_presence("2") == ISOPresenceEnum.Exclusive
    assert resolver.map_bcv_to_presence("1") == ISOPresenceEnum.Conditional

    # Reverse presence mappings
    assert resolver.map_presence_to_bcv("Fixed") == BarCodeValueEnum.BCV_3
    assert resolver.map_presence_to_bcv("Exclusive") == BarCodeValueEnum.BCV_2

    # Coverage factor mappings
    c_min, c_max = resolver.factor_to_coverage("0.55")
    assert c_min == 30 and c_max == 80

    factor = resolver.coverage_to_factor(30, 80)
    assert factor == "055"


def test_eagle_to_iso_transformation():
    assert TEST_EXCEL_PATH.exists()
    wb = openpyxl.load_workbook(TEST_EXCEL_PATH, data_only=True)
    logs = []
    xml_bytes = run_eagle_to_iso_mapping(wb, "CLC-all 3 Classes", log_callback=logs.append)

    assert xml_bytes is not None and len(xml_bytes) > 1000

    # Verify XML well-formedness and structure
    root = ET.fromstring(xml_bytes)
    assert "LC_Legend" in root.tag

    classes = root.findall(".//LC_Class")
    assert len(classes) == 3

    class_codes = [c.findtext("class_map_code") for c in classes]
    assert "C112" in class_codes
    assert "C212" in class_codes
    assert "C322" in class_codes

    props = root.findall(".//LC_Properties")
    assert len(props) > 0


def test_iso_to_eagle_transformation():
    assert TEST_XML_PATH.exists()
    with open(TEST_XML_PATH, "rb") as f:
        xml_data = f.read()

    logs = []
    excel_bytes = run_iso_to_eagle_mapping(xml_data, log_callback=logs.append, use_fuzzy=True)
    assert excel_bytes is not None and len(excel_bytes) > 2000

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    ws = wb.active
    assert ws.title == "EAGLE Barcoding Matrix"

    # Verify class codes in Row 1
    row1_codes = [ws.cell(1, c).value for c in [5, 7, 9]]
    assert "CLC112" in row1_codes or "C112" in row1_codes
    assert "CLC212" in row1_codes or "C212" in row1_codes


def test_roundtrip_transformation():
    wb_in = openpyxl.load_workbook(TEST_EXCEL_PATH, data_only=True)
    xml_bytes = run_eagle_to_iso_mapping(wb_in, "CLC-all 3 Classes", lambda m: None)

    excel_bytes = run_iso_to_eagle_mapping(xml_bytes, lambda m: None, use_fuzzy=True)
    wb_out = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    ws = wb_out.active

    assert ws.cell(1, 5).value in ["C112", "CLC112"]
    assert ws.cell(1, 7).value in ["C212", "CLC212"]
    assert ws.cell(1, 9).value in ["C322", "CLC322"]
